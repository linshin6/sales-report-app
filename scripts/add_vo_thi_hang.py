import sys
import json
import re
import openpyxl
from copy import copy

sys.stdout.reconfigure(encoding='utf-8')

# Target and stores info for Võ Thị Hằng from HNTrinh Form Target PG_T10 2026 NDTRINH (1).xlsb:
HANG_NAME = "Võ Thị Hằng"
HANG_MA_PG = "ST-04A-3177"
HANG_TARGET = 1007529600 # 1,007,529,600 đ
HANG_BHX_STORES = 18
HANG_CVS_STORES = 2
HANG_TOTAL_STORES = 20

HANG_GS25_STORES = [
    {
        "employee_name": HANG_NAME,
        "chain": "GS25",
        "store_code": "CVS_GS25.75.6976",
        "store_address": "GS25 Lê Duẩn - Long Thành",
        "actual": 0
    },
    {
        "employee_name": HANG_NAME,
        "chain": "GS25",
        "store_code": "CVS_GS25.75.6975",
        "store_address": "Siêu thị GS-25 35, Nguyễn Đình Chiểu",
        "actual": 0
    }
]

HANG_BHX_CODES = [
    '4511', '14298', '7011', '4389', '5216', '8221', '30016', '7155',
    'MINI_BHX.75.248', '4754', '6560', '6776', '4752', '6749', '4382',
    '6741', '8108', '4395'
]

print("=== 1. CẬP NHẬT master_data.js ===")
with open('master_data.js', 'r', encoding='utf-8') as f:
    md_content = f.read()

# Check if already added
if HANG_NAME not in md_content:
    # Add to employees
    emp_insert = f''',
    {{
      "name": "{HANG_NAME}",
      "bhx_stores": {HANG_BHX_STORES},
      "target": {HANG_TARGET}
    }}'''
    # Find last employee in employees array
    idx_nhung = md_content.find('"name": "Phạm Thị Kim Nhung"')
    idx_end_nhung = md_content.find('}', idx_nhung)
    md_content = md_content[:idx_end_nhung+1] + emp_insert + md_content[idx_end_nhung+1:]

    # Add 2 GS25 stores to stores array
    stores_insert = f''',
  {{
    "employee_name": "{HANG_NAME}",
    "chain": "GS25",
    "store_code": "CVS_GS25.75.6976",
    "store_address": "GS25 Lê Duẩn - Long Thành",
    "actual": 0
  }},
  {{
    "employee_name": "{HANG_NAME}",
    "chain": "GS25",
    "store_code": "CVS_GS25.75.6975",
    "store_address": "Siêu thị GS-25 35, Nguyễn Đình Chiểu",
    "actual": 0
  }}'''
    # Find last store in stores array: HD3024
    idx_hd = md_content.find('"HD3024"')
    idx_end_hd = md_content.find('}', idx_hd)
    md_content = md_content[:idx_end_hd+1] + stores_insert + md_content[idx_end_hd+1:]

    # Update total_stores_bhx_team: 179 -> 197
    md_content = md_content.replace('"total_stores_bhx_team": 179', '"total_stores_bhx_team": 197')

    with open('master_data.js', 'w', encoding='utf-8') as f:
        f.write(md_content)
    print("✓ master_data.js đã được cập nhật thành công.")
else:
    print("! master_data.js đã có Võ Thị Hằng.")

print("\n=== 2. CẬP NHẬT MasterData.gs ===")
with open('MasterData.gs', 'r', encoding='utf-8') as f:
    gs_content = f.read()

if HANG_NAME not in gs_content:
    emp_insert_gs = f''',
    {{
      "name": "{HANG_NAME}",
      "bhx_stores": {HANG_BHX_STORES},
      "target": {HANG_TARGET}
    }}'''
    idx_nhung = gs_content.find('"name": "Phạm Thị Kim Nhung"')
    idx_end_nhung = gs_content.find('}', idx_nhung)
    gs_content = gs_content[:idx_end_nhung+1] + emp_insert_gs + gs_content[idx_end_nhung+1:]

    stores_insert_gs = f''',
  {{
    "employee_name": "{HANG_NAME}",
    "chain": "GS25",
    "store_code": "CVS_GS25.75.6976",
    "store_address": "GS25 Lê Duẩn - Long Thành",
    "actual": 0
  }},
  {{
    "employee_name": "{HANG_NAME}",
    "chain": "GS25",
    "store_code": "CVS_GS25.75.6975",
    "store_address": "Siêu thị GS-25 35, Nguyễn Đình Chiểu",
    "actual": 0
  }}'''
    idx_hd = gs_content.find('HD3024')
    if idx_hd != -1:
        idx_end_hd = gs_content.find('}', idx_hd)
        gs_content = gs_content[:idx_end_hd+1] + stores_insert_gs + gs_content[idx_end_hd+1:]
    else:
        # find end of stores
        idx_stores_end = gs_content.find('  ],', gs_content.find('"stores":'))
        gs_content = gs_content[:idx_stores_end] + stores_insert_gs + '\n' + gs_content[idx_stores_end:]

    gs_content = gs_content.replace('"total_stores_bhx_team": 179', '"total_stores_bhx_team": 197')

    with open('MasterData.gs', 'w', encoding='utf-8') as f:
        f.write(gs_content)
    print("✓ MasterData.gs đã được cập nhật thành công.")
else:
    print("! MasterData.gs đã có Võ Thị Hằng.")

print("\n=== 3. CẬP NHẬT index.html ===")
with open('index.html', 'r', encoding='utf-8') as f:
    idx_html = f.read()

if HANG_NAME not in idx_html:
    # 1. Add to DEFAULT_MASTER.employees
    nhung_str = "{ name: 'Phạm Thị Kim Nhung', bhx_stores: 20, target: 702989796 }"
    hang_def = f"{nhung_str},\n        {{ name: '{HANG_NAME}', bhx_stores: {HANG_BHX_STORES}, target: {HANG_TARGET} }}"
    idx_html = idx_html.replace(nhung_str, hang_def)
    idx_html = idx_html.replace("total_stores_bhx_team: 179", "total_stores_bhx_team: 197")

    # 2. Add to getInitialStores
    # Find last store in getInitialStores
    last_store_target = '"actual": 0}];'
    if last_store_target in idx_html:
        hang_init_stores = f''', {{"employee_name": "{HANG_NAME}", "chain": "GS25", "store_code": "CVS_GS25.75.6976", "store_address": "GS25 Lê Duẩn - Long Thành", "actual": 0}}, {{"employee_name": "{HANG_NAME}", "chain": "GS25", "store_code": "CVS_GS25.75.6975", "store_address": "Siêu thị GS-25 35, Nguyễn Đình Chiểu", "actual": 0}}'''
        idx_html = idx_html.replace(last_store_target, hang_init_stores + '];')

    with open('index.html', 'w', encoding='utf-8') as f:
        f.write(idx_html)
    print("✓ index.html đã được cập nhật thành công.")
else:
    print("! index.html đã có Võ Thị Hằng.")

print("\n=== 4. CẬP NHẬT team_store_master.json ===")
with open('team_store_master.json', 'r', encoding='utf-8') as f:
    ts_master = json.load(f)

updated_ts = 0
all_hang_codes = set(HANG_BHX_CODES + ['CVS_GS25.75.6976', 'CVS_GS25.75.6975'])
for item in ts_master:
    ms = str(item.get('ma_store') or '').replace('.0', '')
    if ms in all_hang_codes:
        item['ten_pg'] = HANG_NAME
        item['ma_pg'] = HANG_MA_PG
        item['vtcv'] = 'SR'
        item['tinh_trang_lam_viec'] = 'Đang làm'
        updated_ts += 1

with open('team_store_master.json', 'w', encoding='utf-8') as f:
    json.dump(ts_master, f, ensure_ascii=False, indent=2)
print(f"✓ team_store_master.json: Đã cập nhật {updated_ts} cửa hàng cho {HANG_NAME}.")

print("\n=== 5. CẬP NHẬT Team_CamGiang_Report.xlsx ===")
wb = openpyxl.load_workbook('Team_CamGiang_Report.xlsx')

# 5.1 Sheet Config
ws_cfg = wb['Config']
ws_cfg['B10'] = 11 # Số NV trong team
ws_cfg['B11'] = 140 # Tổng CH CVS
print("✓ Config: Số NV = 11, Tổng CH CVS = 140")

# 5.2 Sheet Tiến độ Team
ws_td = wb['Tiến độ Team']
# Check if Võ Thị Hằng already exists
hang_exists_td = False
for r in range(7, 20):
    if ws_td.cell(r, 3).value == HANG_NAME:
        hang_exists_td = True
        break

if not hang_exists_td:
    # Row 17 is currently TỔNG TEAM
    # Insert row at 17
    ws_td.insert_rows(17)
    ws_td.cell(17, 2, 11) # STT
    ws_td.cell(17, 3, HANG_NAME) # Nhân Viên
    ws_td.cell(17, 4, HANG_TARGET) # Target
    bhx_cd = round(18 * 19954649) # BHX (cố định)
    ws_td.cell(17, 5, bhx_cd)
    ws_td.cell(17, 6, 0) # GS25
    ws_td.cell(17, 7, 0) # 7-Eleven
    ws_td.cell(17, 8, 0) # FamilyMart
    ws_td.cell(17, 9, 0) # Circle K
    ws_td.cell(17, 10, 0) # Hoàng Đức
    ws_td.cell(17, 11, 0) # WinMart+
    ws_td.cell(17, 12, bhx_cd) # Tổng TH
    ws_td.cell(17, 13, bhx_cd / HANG_TARGET) # % Đạt

    # Copy style from row 16
    for c in range(1, 14):
        src = ws_td.cell(16, c)
        dst = ws_td.cell(17, c)
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = copy(src.number_format)
            dst.alignment = copy(src.alignment)

    # Row 18 is now TỔNG TEAM
    for col_idx, col_letter in enumerate(['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L'], start=4):
        ws_td.cell(18, col_idx).value = f"=SUM({col_letter}7:{col_letter}17)"
    ws_td.cell(18, 13).value = "=L18/D18"
    print("✓ Tiến độ Team: Đã thêm hàng 17 cho Võ Thị Hằng.")

# 5.3 Sheet Dashboard
ws_db = wb['Dashboard']
hang_exists_db = False
for r in range(13, 25):
    if ws_db.cell(r, 3).value == HANG_NAME:
        hang_exists_db = True
        break

if not hang_exists_db:
    r_hang = 23
    ws_db.cell(r_hang, 2, 11) # Hạng
    ws_db.cell(r_hang, 3, HANG_NAME) # Nhân Viên
    ws_db.cell(r_hang, 4, HANG_TARGET) # Target
    bhx_cd = round(18 * 19954649)
    ws_db.cell(r_hang, 5, bhx_cd) # BHX
    ws_db.cell(r_hang, 6, 0) # CVS
    ws_db.cell(r_hang, 7, bhx_cd) # Tổng TH
    ws_db.cell(r_hang, 8, bhx_cd / HANG_TARGET) # % Đạt

    # Copy style from row 22
    for c in range(1, 9):
        src = ws_db.cell(22, c)
        dst = ws_db.cell(r_hang, c)
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = copy(src.number_format)
            dst.alignment = copy(src.alignment)

    # Update summary row 7 in Dashboard
    # Target Team, Thực Hiện BHX, Thực Hiện CVS, Tổng Thực Hiện, Thiếu, % Đạt Team
    tot_tgt_team = 10826339729 + HANG_TARGET
    tot_bhx_team = 3551927443 + bhx_cd
    tot_cvs_team = 359247028
    tot_th_team = tot_bhx_team + tot_cvs_team
    ws_db.cell(7, 1, tot_tgt_team)
    ws_db.cell(7, 2, tot_bhx_team)
    ws_db.cell(7, 3, tot_cvs_team)
    ws_db.cell(7, 4, tot_th_team)
    ws_db.cell(7, 5, tot_tgt_team - tot_th_team)
    ws_db.cell(7, 6, tot_th_team / tot_tgt_team)
    print("✓ Dashboard: Đã thêm hàng 23 cho Võ Thị Hằng và cập nhật tổng quan.")

# 5.4 Sheet Chi tiết CVS
ws_cvs = wb['Chi tiết CVS']
hang_exists_cvs = False
for r in range(1, ws_cvs.max_row+1):
    if HANG_NAME in str(ws_cvs.cell(r, 1).value or '') or HANG_NAME in str(ws_cvs.cell(r, 2).value or ''):
        hang_exists_cvs = True
        break

if not hang_exists_cvs:
    # Find total row at bottom
    total_r = None
    for r in range(ws_cvs.max_row - 10, ws_cvs.max_row + 1):
        if 'TỔNG CVS TOÀN TEAM' in str(ws_cvs.cell(r, 1).value or ''):
            total_r = r
            break
    
    if total_r:
        # Insert 4 rows before total_r
        # Row 1: Section header
        # Row 2: Store 1 (GS25 Lê Duẩn)
        # Row 3: Store 2 (GS25 35 Nguyễn Đình Chiểu)
        # Row 4: Empty separator
        ws_cvs.insert_rows(total_r, 4)
        
        # Section header
        ws_cvs.cell(total_r, 1, f"👤 {HANG_NAME} (2 cửa hàng CVS)")
        ws_cvs.cell(total_r, 5, "TỔNG CVS:")
        ws_cvs.cell(total_r, 6, 0)
        
        # Store 1
        ws_cvs.cell(total_r + 1, 1, 139)
        ws_cvs.cell(total_r + 1, 2, HANG_NAME)
        ws_cvs.cell(total_r + 1, 3, "GS25")
        ws_cvs.cell(total_r + 1, 4, "CVS_GS25.75.6976")
        ws_cvs.cell(total_r + 1, 5, "GS25 Lê Duẩn - Long Thành")
        ws_cvs.cell(total_r + 1, 6, 0)
        
        # Store 2
        ws_cvs.cell(total_r + 2, 1, 140)
        ws_cvs.cell(total_r + 2, 2, HANG_NAME)
        ws_cvs.cell(total_r + 2, 3, "GS25")
        ws_cvs.cell(total_r + 2, 4, "CVS_GS25.75.6975")
        ws_cvs.cell(total_r + 2, 5, "Siêu thị GS-25 35, Nguyễn Đình Chiểu")
        ws_cvs.cell(total_r + 2, 6, 0)
        
        # Styling
        # Copy header style from row 156 (Pham Thi Kim Nhung header)
        # Find Pham Thi Kim Nhung header
        nhung_hdr_r = None
        for r in range(1, total_r):
            if 'Phạm Thị Kim Nhung' in str(ws_cvs.cell(r, 1).value or ''):
                nhung_hdr_r = r
                break
        
        if nhung_hdr_r:
            for c in range(1, 7):
                src = ws_cvs.cell(nhung_hdr_r, c)
                dst = ws_cvs.cell(total_r, c)
                if src.has_style:
                    dst.font = copy(src.font)
                    dst.border = copy(src.border)
                    dst.fill = copy(src.fill)
                    dst.alignment = copy(src.alignment)
            
            # Store styles from nhung_hdr_r + 1
            src_st = ws_cvs.cell(nhung_hdr_r + 1, 1)
            for offset in [1, 2]:
                for c in range(1, 7):
                    s_c = ws_cvs.cell(nhung_hdr_r + 1, c)
                    d_c = ws_cvs.cell(total_r + offset, c)
                    if s_c.has_style:
                        d_c.font = copy(s_c.font)
                        d_c.border = copy(s_c.border)
                        d_c.fill = copy(s_c.fill)
                        d_c.number_format = copy(s_c.number_format)
                        d_c.alignment = copy(s_c.alignment)
        
        print("✓ Chi tiết CVS: Đã thêm mục cho Võ Thị Hằng với 2 cửa hàng GS25.")

wb.save('Team_CamGiang_Report.xlsx')
print("✓ Team_CamGiang_Report.xlsx đã được lưu thành công.")
